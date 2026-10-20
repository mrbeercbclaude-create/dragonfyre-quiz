import json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
G=json.load(open('guide.json',encoding='utf-8'))
FONT='Leelawadee UI'
INK=RGBColor(0x2b,0x21,0x18);INK2=RGBColor(0x5c,0x4d,0x3f);MUTED=RGBColor(0x8a,0x7a,0x69);CLAY=RGBColor(0xb5,0x55,0x2f)
GOOD=RGBColor(0x4f,0x8a,0x3c);KEY=RGBColor(0x4f,0x73,0x93);GOLD=RGBColor(0xc8,0x90,0x2e);GRAY=RGBColor(0xa9,0x9a,0x86);WHITE=RGBColor(255,255,255)
BM=[0]
def fontify(run,size=None,bold=None,color=None,italic=None):
    rPr=run._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rPr.insert(0,rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a),FONT)
    if bold: run.bold=True; rPr.append(OxmlElement('w:bCs'))
    if italic: run.italic=True; rPr.append(OxmlElement('w:iCs'))
    if color is not None: run.font.color.rgb=color
    if size:
        run.font.size=Pt(size); s=OxmlElement('w:szCs'); s.set(qn('w:val'),str(int(size*2))); rPr.append(s)
    return run
def para(doc,text='',size=11,bold=False,color=INK,italic=False,align=None,after=4,before=0):
    p=doc.add_paragraph(); f=p.paragraph_format; f.space_after=Pt(after); f.space_before=Pt(before); f.line_spacing=1.15
    if align is not None: p.alignment=align
    if text: fontify(p.add_run(text),size,bold,color,italic)
    return p
def shade_el(el,hexcol):
    sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexcol); el.append(sh)
def keep_next(p): p.paragraph_format.keep_with_next=True
def pagebreak(doc): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def heading(doc,text,size,color=CLAY,before=10,after=4,align=None):
    p=para(doc,'',before=before,after=after,align=align)
    if text: fontify(p.add_run(text),size,True,color)
    keep_next(p); return p
def field(p,code,size=9):
    r=p.add_run(); fontify(r,size,color=MUTED)
    for t,c in (('begin',None),(None,code),('end',None)):
        if t: e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),t)
        else: e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=c
        r._element.append(e)
def table(doc,rows,widths,head):
    t=doc.add_table(rows=0,cols=len(widths)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,row in enumerate([head]+rows):
        cells=t.add_row().cells
        for j,(c,w) in enumerate(zip(row,widths)):
            cells[j].width=Cm(w); p=cells[j].paragraphs[0]; p.paragraph_format.space_after=Pt(1)
            fontify(p.add_run(c),10 if i==0 else 9.5,i==0 or j==0,WHITE if i==0 else (INK if j==0 else INK2))
            if i==0: shade_el(cells[j]._element.get_or_add_tcPr(),'5C4D3F')
            elif i%2==0: shade_el(cells[j]._element.get_or_add_tcPr(),'F6EFE2')
    return t
STAR=lambda n:'★'*n+'☆'*(3-n)
SYM={'+':('✚',GOOD),'-':('✖',CLAY),'*':('⌨',KEY),'=':('ⓘ',GRAY)}
MODS=["Cisco","Origins++","Greek","Fairytale","Magic","Mythic","Apocalypse","Chidori","Animal","Revised","SF","Kitsune"]
MODNAME={"Cisco":"Cisco's RPG Origins (เผ่าหลักของแพ็ก)","Origins++":"Origins++","Greek":"Greek Mythology Origins","Fairytale":"Fairytale Origins","Magic":"Magic Origins","Mythic":"Mythic Origins",
 "Apocalypse":"Apocalypse Origins","Chidori":"Chidori Origins","Animal":"Animal Origins","Revised":"Origins Revised","SF":"Simply Fantastical Origins","Kitsune":"Origins Kitsune"}
MODINTRO={"Cisco":"เผ่าดั้งเดิมของ Dragonfyre ออกแบบมาให้บาลานซ์กับแพ็กโดยตรง มีเผ่าสายเวทที่ผูกกับ Iron's Spells (Justicar, Dark Mage, Phoenix Knight) เหมาะกับคนที่อยากได้ของที่เทสต์มาแล้ว",
 "Origins++":"ม็อดใหญ่ที่สุด 120 กว่าเผ่า ตั้งแต่สัตว์ หุ่นยนต์ ปีศาจ ไปจนถึงสิ่งจากอีกมิติ หลายเผ่ามีหลอดพลังและระบบเฉพาะตัว ★★★ เยอะ อ่านพลังให้ครบก่อนเลือก",
 "Greek":"ลูกของเทพกรีก 13 องค์ แต่ละเผ่ามีสกิลกดเด่นๆ และข้อเสียตามตำนาน เช่น ลูกซุสอ่อนแรงใต้ดิน ลูกโพไซดอนแห้งในทะเลทราย",
 "Fairytale":"ตัวละครจากนิทาน หมาป่า นางเงือก มิโนทอร์ ซินเดอเรลล่า เล่นสนุกและมีธีมชัด บางพลังใช้ไม่ได้ในเซิร์ฟเรา",
 "Magic":"มังกร หุ่นไซบอร์ก นักปลุกผี หยินหยาง ฯลฯ มีเผ่าสายช่วยทีมแรงๆ อย่าง Magical Fly",
 "Mythic":"สิ่งมีชีวิตในตำนาน 5 เผ่า ฟีนิกซ์ที่ตายแล้วฟื้น วาลคิรีบินได้ จินน์ทราย จิ้งจอก และดรูอิด",
 "Apocalypse":"เผ่าที่ลึกที่สุดในเซิร์ฟ: นักพันธุศาสตร์ฉีดยีนมอน ดาบผี Hexblade เห็ดที่วิวัฒนาการได้ 9 ร่าง และคริสตัลวาร์ป เหมาะกับคนชอบระบบ",
 "Chidori":"ลูกผสม 4 ธาตุ ลม น้ำ ความมืด แผ่นดิน มีกระเป๋าพิเศษติดตัวเกือบทุกเผ่า",
 "Animal":"สัตว์ 6 ชนิด กิ้งก่าพรางตัว จิงโจ้ ลิง เพนกวิน สัตว์เลื้อยคลาน และนกที่วิ่งเร็วที่สุด",
 "Revised":"เผ่าแร่และธาตุ 9 แบบ ทองแดง ปะการัง กระดูก ราก ร่างแสง ฯลฯ มีระบบแปลกๆ ให้เล่น",
 "SF":"เผ่าแฟนตาซีคลาสสิก เอลฟ์ คนแคระ ออร์ค ทีฟลิง เล่นง่ายทุกเผ่า (★) เหมาะกับมือใหม่ (Goblin โหลดไม่ขึ้นในเซิร์ฟเรา)",
 "Kitsune":"จิ้งจอกโบราณ เลือกสีแดงหรือขาวได้ พลังเหมือนกัน"}
doc=Document()
sec=doc.sections[0]; sec.page_height=Cm(29.7); sec.page_width=Cm(21.0)
sec.left_margin=sec.right_margin=Cm(1.9); sec.top_margin=sec.bottom_margin=Cm(1.7)
st=doc.styles['Normal']; st.font.name=FONT; st.font.size=Pt(11)
st.element.rPr.rFonts.set(qn('w:cs'),FONT); st.element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; fontify(fp.add_run('Dragonfyre Codex · ZenixSg · หน้า '),9,color=MUTED); field(fp,'PAGE')
C=WD_ALIGN_PARAGRAPH.CENTER
# ===== cover =====
para(doc,'',after=70)
para(doc,"ZENIXSG  ·  CISCO'S DRAGONFYRE",11,True,GOLD,align=C,after=6)
para(doc,'DRAGONFYRE CODEX',40,True,INK,align=C,after=4)
para(doc,'คู่มือเผ่า อาชีพ และพรจากเทพ ฉบับภาษาไทย',18,False,CLAY,align=C,after=30)
nP=sum(len(o['p']) for o in G['O'])
t=doc.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER
for c,(b,s) in zip(t.rows[0].cells,[(len(G['O']),'เผ่า'),(len(G['C']),'อาชีพ'),(len(G['B']),'พรจากเทพ'),(f'{nP:,}','พลังที่แปลแล้ว')]):
    c.width=Cm(3.8); shade_el(c._element.get_or_add_tcPr(),'F7D9A6'); p=c.paragraphs[0]; p.alignment=C; fontify(p.add_run(str(b)),24,True,INK)
    p2=c.add_paragraph(); p2.alignment=C; fontify(p2.add_run(s),10,color=INK2)
para(doc,'',after=34)
para(doc,'แปลจากไฟล์ mod ที่เซิร์ฟเราใช้จริง: Dragonfyre 3C Beta + ม็อดเผ่าเสริม 11 ตัว + พรกรีกที่เซิร์ฟทำเอง',10.5,color=INK2,align=C,after=2)
para(doc,'ทำแบบทดสอบหาบิลด์: mrbeercbclaude-create.github.io/dragonfyre-quiz',10.5,True,KEY,align=C,after=2)
para(doc,'ค้นหาแบบออนไลน์: mrbeercbclaude-create.github.io/dragonfyre-quiz/guide.html',10.5,True,KEY,align=C,after=2)
para(doc,'ฉบับ 20 ตุลาคม 2026',10,color=MUTED,align=C)
pagebreak(doc)
# ===== how to read =====
heading(doc,'อ่านก่อนเลือก',20,before=0)
for b,tx in [('เลือก 3 อย่างตอนเข้าเซิร์ฟครั้งแรก: ','เผ่า (Origin) แล้วอาชีพ (Class) แล้วพรจากเทพ (Divine Blessing) สามอย่างนี้รวมกันคือ "บิลด์" ของคุณ'),
 ('เลือกผิดแก้ได้: ','quest แรกแจก Orb of Origin ฟรี 1 ลูก และหาเพิ่มได้จากกล่องสุ่มในร้าน quest'),
 ('ตั้งปุ่มก่อนเล่น: ','Options → Controls ตั้ง View Origin, Active Power (Primary), Active Power (Secondary) และ Save/Load Toolbar Activator ถ้าเผ่าใช้ ไม่ตั้ง = ใช้สกิลกดไม่ได้'),
 ('ดาวความแรง: ','★ เล่นง่าย · ★★ มีจุดอ่อนต้องระวัง · ★★★ เปลี่ยนวิธีเล่นทั้งเกม แรงแต่ราคาแพง'),
 ('ยังเลือกไม่ได้? ','ทำแบบทดสอบ 30 ข้อก่อน ได้บิลด์ 10 ชุดที่เข้ากับคุณ แล้วกลับมาเปิดหาเผ่านั้นในคู่มือนี้'),
 ('เผ่าจาก Cisco ','ใช้คำอธิบายจากคู่มือเดิมของเซิร์ฟ ส่วนเผ่าจากม็อดเสริมแปลทีละพลังจากไฟล์ jar')]:
    p=para(doc,after=5); fontify(p.add_run('•  '+b),11,True,INK); fontify(p.add_run(tx),11,color=INK2)
heading(doc,'สัญลักษณ์ในคู่มือ',14,before=12)
for k,lab in [('+','ข้อดี'),('-','ข้อเสีย'),('*','สกิลกดใช้ (ต้องตั้งปุ่ม)'),('=','ข้อมูลหรือระบบพิเศษของเผ่า')]:
    p=para(doc,after=2); s,c=SYM[k]; fontify(p.add_run(s+'  '),12,True,c); fontify(p.add_run(lab),11,color=INK2)
heading(doc,'สารบัญ',14,before=12)
for name,cnt in [('พรจากเทพ',len(G['B'])),('อาชีพ',len(G['C']))]+[(MODNAME[m],sum(1 for o in G['O'] if o['m']==m)) for m in MODS]:
    p=para(doc,after=1); fontify(p.add_run(name),11,color=INK); fontify(p.add_run(f'  ·  {cnt}'),10,color=MUTED)
pagebreak(doc)
# ===== blessings =====
heading(doc,'พรจากเทพ (Divine Blessing)',22,before=0)
para(doc,'พรของ Cisco มี 2 ขั้น: ได้พลังเล็กทันที และพลังใหญ่หลังฆ่า Ender Dragon · พรกรีก 10 อันเซิร์ฟเราทำเพิ่ม เป็นแบบข้อดี 2 อย่าง + ข้อเสีย 1 อย่าง ไม่มีขั้นหลังฆ่ามังกร · Ares, Neptune และ No Blessing ของ Cisco โหลดไม่ขึ้นในเซิร์ฟเรา',10,color=INK2,after=8)
cis=[b for b in G['B'] if b['m']=='Cisco']; gr=[b for b in G['B'] if b['m']!='Cisco']
heading(doc,'พรของ Cisco',14,before=4)
table(doc,[[b['n'],b.get('now',''),b.get('late','')] for b in cis],[3.0,6.6,7.6],['พร','ได้ทันที','หลังฆ่า Ender Dragon'])
heading(doc,'พรกรีกของเซิร์ฟเรา',14,before=12)
table(doc,[[b['n'],b['d'].split(' · ')[0],(b['d'].split(' · ')[1] if ' · ' in b['d'] else '-')] for b in gr],[3.6,8.4,5.2],['พร','ข้อดี','ข้อเสีย'])
pagebreak(doc)
# ===== classes =====
heading(doc,'อาชีพ (Class)',22,before=0)
para(doc,'อาชีพของ Cisco ไม่มีข้อเสีย เลือกตามงานที่ทำบ่อยที่สุด · อาชีพจาก Apocalypse บางอันเลือกไม่ได้กับบางเผ่า',10,color=INK2,after=8)
table(doc,[[c['n']+('' if c['m']=='Cisco' else '  (Apocalypse)'),(c.get('full') or c['d'])+(f"  · เลือกไม่ได้ถ้าเผ่าเป็น {', '.join(c['no'])}" if c.get('no') else '')] for c in G['C']],[4.4,12.8],['อาชีพ','ได้อะไร'])
# ===== origins =====
for m in MODS:
    os_=sorted([o for o in G['O'] if o['m']==m],key=lambda o:(o['i'],o['n']))
    if not os_: continue
    pagebreak(doc)
    heading(doc,MODNAME[m],24,before=0,after=2)
    para(doc,f'{len(os_)} เผ่า · เรียงจากเล่นง่ายไปยาก',11,True,GOLD,after=4)
    pi=para(doc,MODINTRO[m],10.5,color=INK2,after=10); shade_el(pi._element.get_or_add_pPr(),'F6EFE2')
    for o in os_:
        p=heading(doc,'',15,before=12,after=1); fontify(p.add_run(o['n']),15,True,INK); fontify(p.add_run('   '+STAR(o['i'])),12,False,GOLD)
        p=para(doc,o['d'],10.5,color=INK2,italic=True,after=2); keep_next(p)
        tags=[G['P'].get(t,t) for t in o['tg'][:5]]; warn=[G['F'][f] for f in o['f']]
        if tags or warn:
            p=para(doc,after=3); keep_next(p)
            if tags: fontify(p.add_run('เด่น: '+' · '.join(tags)),9.5,True,GOOD)
            if warn: fontify(p.add_run(('     ' if tags else '')+'ระวัง: '+' · '.join(warn)),9.5,True,CLAY)
        for lab,v,col in [('วิวัฒนาการ: ',o['ev'],KEY),('เคล็ดลับ: ',o['x'],INK),('ในเซิร์ฟเรา: ',o['br'],CLAY)]:
            if v: p=para(doc,after=2); fontify(p.add_run(lab),10,True,col); fontify(p.add_run(v),10,color=INK2); keep_next(p)
        for k in ('*','+','=','-'):
            for x in o['p']:
                if x['t']!=k: continue
                p=para(doc,after=1); p.paragraph_format.left_indent=Cm(0.6); p.paragraph_format.first_line_indent=Cm(-0.6)
                s,c=SYM[k]; fontify(p.add_run(s+'  '),10,True,c)
                if x['n']: fontify(p.add_run(x['n']+' — '),10,True,INK)
                fontify(p.add_run(x['d']),10,color=INK2)
pagebreak(doc)
heading(doc,'แล้วเจอกันในเซิร์ฟ',22,before=120,after=6,align=C)
para(doc,'ถ้าเจอพลังที่แปลผิด หรือเผ่าที่ใช้ไม่ได้จริง บอก PonderingBaby ได้เลย',11,color=INK2,align=C)
out='C:/Users/mrbee/Downloads/dragonfyre-quiz/Dragonfyre-Codex.docx'
doc.core_properties.title='Dragonfyre Codex'; doc.core_properties.author='ZenixSg'
doc.save(out); print('saved',out)
